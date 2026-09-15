"""Build auditable XLSX and UTF-8 CSV exports from research inputs (no network)."""
from pathlib import Path
import csv, json, math, zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from research_data import *
from research_profiles import *
from research_schema import SCHEMAS
ROOT=Path(__file__).resolve().parent
OUT=ROOT/'deliverables'; OUT.mkdir(exist_ok=True)
CSV=OUT/'csv'; CSV.mkdir(exist_ok=True)
H='HYPOTHESIS: '; A='ASSUMPTION: '; U='DATA NOT VERIFIED'
sheets={k:[v] for k,v in SCHEMAS.items()}
def add(k, row):
    assert len(row)==len(sheets[k][0]), (k,len(row),len(sheets[k][0]))
    sheets[k].append(row)
def supplement(k,headers): sheets[k]=[headers.split('|')]
def rec(o):
    if 'REJECT' in o['risk']: return 'REJECT'
    if o['id']=='SAAS-001': return 'BUILD NOW'
    if o['id'] in ['SAAS-002','SAAS-003','SAAS-004','SAAS-005']: return 'STRONG ALTERNATIVE'
    if o['id']=='SAAS-033': return 'WATCH'
    return 'VALIDATE' if o['score']>=68 else 'WATCH'
def risknum(o):
    if rec(o)=='REJECT': return 8
    if o['id']=='SAAS-001': return 6
    return 7 if o['id'] in ['SAAS-005','SAAS-008','SAAS-010','SAAS-013','SAAS-019','SAAS-033'] else 6

def channels(o): return CHANNELS.get(o['id'][-3:],[max(2,o['f']['CAC']),3,3,3,5,5,3,5,6,6])
def price(o):
    p=PROFILES.get(o['id']); return A+f"MAD{p['mad']:,.0f}/mo or USD{p['usd']:,.0f}/mo internationally" if p else U

for o in OPPS:
    f=o['f']; c=channels(o); p=PROFILES.get(o['id']); recommendation=rec(o)
    if recommendation=='BUILD NOW': recommendation+=' — paid validation/concierge only; gate full MVP'
    add('01_MASTER_RANKING',[o['id'],o['rank'],o['name'],o['cat'],o['market'],o['country'],o['icp'],o['pain'],f['Demand'],f['Pain'],f['WTP'],f['Recurring'],f['HighTicket'],max(c[:3]),f['CAC'],f['Retention'],f['Competition'],f['MVPEase'],f['AI'],f['Global'],f['Morocco'],f['ROI'],f['Moat'],o['score'],'Vertical workflow' if o['id'] in TOP15 else 'Category screen',recommendation])
    add('02_GLOBAL_MARKETS',[o['id'],o['rank'],o['cat'],p['product'] if p else o['name'],o['icp'],o['market'],f"Category spending/capability proxy [{o['sources']}]; exact niche WTP UNKNOWN",'UNKNOWN niche growth; no invented CAGR',o['pain'],H+str(f['WTP'])+'/10',H+'Monthly or annual subscription',9,f['HighTicket'],c[0],c[1],c[2],c[7],c[4],f['Competition'],f['MVPEase'],f['Global'],o['score'],H+'Frequent workflow with identifiable owner; test incremental value over incumbent',o['risk']])
    add('14_RISK_ANALYSIS',[o['id'],o['name'],'Niche buyer count and paid demand UNKNOWN','Workflow/data quality and adoption',o['risk'],'CAC modeled, never observed','Logo churn unobserved; bundling may remove value','Imports, access controls and integration drift','Local legal review; CNDP/GDPR where applicable','No core WhatsApp collection dependency for001; no scraped/private APIs','Translation alone is not differentiation','Pilot receipts and prepayment; monitor own overdue AR','Founder selling and onboarding concentration',risknum(o),'Paid pilots; 1 ICP/1 export; exception-specific ROI; human approvals; backups; exportable data'])
    add('15_FINAL_DECISION',[o['id'],o['rank'],o['name'],o['market'],o['icp'],price(o),f['HighTicket'],f['Demand'],f['Pain'],f['WTP'],f['CAC'],max(c[:3]),11-f['CAC'],f['Retention'],f['MVPEase'],f['Competition'],f['Moat'],f['Morocco'],f['Global'],f['AI'],f['ROI'],risknum(o),o['score'],'1: paid validation, not full build' if o['id']=='SAAS-001' else ('Gated alternative' if rec(o)=='STRONG ALTERNATIVE' else 'Do not build yet'),rec(o)])

for o in [x for x in OPPS if x['id'] in TOP15]:
    i=o['id']; f=o['f']; p=PROFILES[i]; c=channels(o)
    add('03_MOROCCO_MARKETS',[i,o['rank'],o['name'],o['cat'],o['icp'],o['icp'].split(';')[1].strip(),o['pain'],H+'Spreadsheet + existing industry suite + email/phone; WhatsApp usage verify',H+'French onboarding, MAD invoices, local exports and exception-specific workflow','Important for customer-facing messages; role-dependent','French-first admin hypothesis','Darija onboarding/support; voice AI not MVP','Opted-in service/sales only;001 collections excluded [S04]','Bank transfer pilot; PSP/acquirer eligibility validate',H+'Company-level subscription; volume caps',A+f"MAD{p['mad']*.7:.0f}–{p['mad']*1.5:.0f}/month",f['HighTicket'],'Odoo/Sage integrators; schools E-Schools/Minassa; clinics DabaDoc; exhaustive list UNKNOWN','See08_COMPETITORS; do not assume poor French localization',max(c[0],c[1]),f['CAC'],f['Retention'],'Qualified Moroccan establishment count UNKNOWN; OMPIC/TIC not TAM [S05;S32]',o['score']])
    add('04_MOROCCO_GLOBAL',[i,o['rank'],p['product'],o['icp']+' in Morocco',H+'Founder access, bilingual onboarding, local workflow/export knowledge',H+o['pain'],'Senegal or Côte d’Ivoire: validate one country','France then Belgium: EU compliance + local connectors','Canada then GCC; not simultaneous',o['icp'],o['pain'],'See20_EXPANSION_PLAYBOOK; language, PSP, privacy, invoice fields, support and connectors vary','See08_COMPETITORS; functionality not unique',f"MAD{p['mad']:.0f}/mo (A)",f"EUR{p['eur']:.0f} / USD{p['usd']:.0f}/mo (A)",7,f['Global'],o['score']])
    if f['HighTicket']>=7:
        hi=max(599,p['usd']*2)
        add('05_HIGH_TICKET',[i,o['rank'],p['product'],o['cat'],o['icp'],p['buyer'],A+'USD'+p['rev']+'/year',f['Pain'],A+'USD'+p['cost']+'/month',H+'Require documented monthly benefit >=3× subscription; no recovery guarantee',A+f"USD{hi:.0f}–{hi*2:.0f}",A+f"USD{hi*12:.0f}–{hi*24:.0f}",f['HighTicket'],A+('4–12 weeks' if i in ['SAAS-005','SAAS-010','SAAS-013'] else '2–8 weeks'),11-f['CAC'],c[0],c[2],c[4],f['CAC'],f['Retention'],H+'Moderate: historical exceptions, policies and connector; preserve exportability',f['Competition'],o['score']])
    add('06_ADVERTISING',[i,o['icp'],U+'; keyword themes: '+p['keyword'],'Commercial workflow intent plausible; CPC/volume UNKNOWN',*c,p['channel'],A+f"USD{p['caclo']:.0f}–{p['cachi']:.0f} international all-in blended",H+'Buyer recognizes costly workflow; test booked-demo-to-paid not form leads','Low search volume, suite incumbents, jobseekers/consumer clicks; ads not proven profitable'])
    add('07_CUSTOMER_PROFILE',[i,o['icp'],o['cat'],o['icp'].split(';')[1].strip(),o['country']+' beachhead; wider regions conditional',p['buyer'].split(';')[0],p['buyer'],o['icp'].split(';')[1].strip(),A+'USD'+p['rev']+'/year; Moroccan qualifying revenue differs',o['pain'],'Daily/weekly; school finance monthly and term cycle',H+p['workflow'],H+'Excel/email + '+', '.join(x[0] for x in COMPETITORS.get(i[-3:],[])[:3]),A+'USD'+p['cost']+'/month',H+'Fewer exceptions; shorter cycle time; measurable labor/contribution benefit',p['trigger'],p['objection'],A+f"MAD{p['mad']:.0f}; EUR{p['eur']:.0f}; USD{p['usd']:.0f} per month",p['keyword'],H+'Google, industry association directories, supplier groups and LinkedIn; validate audience reach',p['channel']])
    mobile='Responsive mobile/PWA; camera upload' if i in ['SAAS-002','SAAS-003','SAAS-004','SAAS-010','SAAS-011'] else 'Responsive web; no native app'
    add('09_PRODUCT',[i,p['product'],'Resolve '+o['pain'].lower(),o['icp'],o['pain'],'Narrow exception overlay; do not replace system of record',p['workflow'],*p['features'],p['ai'],'Deterministic queues/reminders; human approval before irreversible actions',p['integrations'],'Exceptions by owner, amount, age and next action','Weekly baseline vs observed time and resolved exceptions','Email/in-app; approved SMS optional; no WhatsApp collections','Admin; operator; reviewer; read-only accountant/client with scoped links',mobile,'Multi-tenant web with RBAC, audit, backups, import/export','Internal API; CSV first; public API only after demand',p['complex'][0],A+p['weeks']+' weeks;2 engineers; compliance may delay launch'])
    add('11_VALIDATION',[i,o['icp'],H+'Unrelated buyers pay target price for observable workflow improvement','NO — not created',0,0,0,0,0,0,0,0,'NOT RUN',U,p['objection']+' (anticipated, not observed)',U,'NOT TESTED','RETEST','Desk research only. All funnel counts are actual work completed=0; no fabricated traction.'])
    # Pricing: all tiers in three currencies; displayed ARPU is proposed professional baseline, not empirical mix.
    for market,curr,key,fx in [('Morocco','MAD','mad',.1),('Europe','EUR','eur',1.1),('USA / International','USD','usd',1)]:
        for plan,mult in [('Free strategy',0),('Trial strategy',0),('Starter',.6),('Professional',1),('Business',2),('Enterprise',4),('High-ticket / custom',6)]:
            monthly=round(p[key]*mult/10)*10 if mult else 0
            if plan=='Professional': monthly=p[key]
            arpu=p[key]; usd_arpu=arpu*fx
            ca_lo,ca_hi=(250,900) if market=='Morocco' else (p['caclo'],p['cachi'])
            gmlo,gmhi=p['gmlo']/100,p['gmhi']/100; lo,hi=p['churnlo']/100,p['churnhi']/100
            ltvlo=usd_arpu*gmlo*min(24,1/hi); ltvhi=usd_arpu*gmhi*min(24,1/lo)
            desc='No permanent free production plan; redacted-data audit' if plan=='Free strategy' else ('14-day sandbox;30-day paid concierge pilot at Pro; no fake free delivery' if plan=='Trial strategy' else 'Proposed subscription; annual 10% discount; custom tiers are floors, not quotes')
            if i=='SAAS-001':
                desc+='; Pilot/Professional:1 entity,3 users,500 open invoices,1 agreed CSV format; Starter:1 user/150 open invoices; Business:10 users/2000 open invoices; Enterprise/custom:volume and support negotiated. No unlimited concierge work.'
            add('10_PRICING_UNIT_ECONOMICS',[i,market,plan,f'{curr}{monthly:.0f}' if mult else 'No ongoing free plan / sandbox only',f'{curr}{monthly*12*.9:.0f}' if mult else 'N/A',o['icp'],f'{curr}{arpu:.0f} professional baseline; mixed ARPU UNKNOWN',f'{p["gmlo"]:.0f}–{p["gmhi"]:.0f}',f'USD{ca_lo:.0f}–{ca_hi:.0f}',f'{p["churnlo"]}–{p["churnhi"]}',f'USD{ltvlo:.0f}–{ltvhi:.0f} (24-month capped GP proxy)',f'{ltvlo/ca_hi:.1f}–{ltvhi/ca_lo:.1f}',f'{ca_lo/(usd_arpu*gmhi):.1f}–{ca_hi/(usd_arpu*gmlo):.1f}',math.ceil(10000/usd_arpu),math.ceil(50000/usd_arpu),math.ceil(100000/usd_arpu),'LOW — all economics modeled',A+desc+'; FX: MAD10/USD1, EUR1/USD1.10 not spot rates; taxes/usage/implementation extra; CAC and LTV USD; early concierge GM may be30–60%; no claim of measured LTV'])

for sid, comps in COMPETITORS.items():
    i='SAAS-'+sid; p=PROFILES[i]
    for name,web,country,position,pricing in comps:
        free={'Upflow':'Product page advertises free AR analytics; collection plan price UNKNOWN [S35]','Odoo':'One App Free subject to scope [S08]','MaintainX':'Basic free [S23]','Eduka':'Express <=100 pupils, finance scope confirm [S21]','Jobber':'No permanent free plan stated;14-day trial [S09]','Cliniko':'30-day trial; no general free plan [S20]','AutoRFP.ai':'Paid30-day qualifying money-back, not free [S17]'}.get(name,U)
        ai={'Upflow':'Vendor-described AR insights, collection and cash-matching agents [S35]','Conexiom':'Vendor-described order extraction/ERP checks [S15]','MaintainX':'Procedure generation/anomaly tools [S23]','AutoRFP.ai':'Cited AI response automation [S17]','Productive':'AI assistant/time tracking/reporting [S18]','Jobber':'Receptionist and quote drafting [S09]'}.get(name,U)
        sentiment={'Chaser':'Indexed G2 praise for reminder automation; selected summary only [S31]','MaintainX':'Capterra mobile praise; some connectivity complaints; selected summary [S24]'}.get(name,U+'; no representative review sample collected')
        add('08_COMPETITORS',[i,name,'https://'+web,country+' (origin hypothesis; current HQ not verified)' if country!=U else U,position,pricing,free,position,ai,'Established workflow coverage; incumbent trust (analyst assessment)','Wedge hypothesis only: greater suite breadth may raise setup burden; no blanket localization weakness claimed',sentiment,'Direct/adjacent comparison; capabilities not fully hands-on tested',U+'; no Meta/Google ad-library audit',U+'; no traffic/domain-authority measurement','Existing integrations, suite coverage and switching inertia (hypothesis)',H+p['product']+': '+p['features'][1]+' + hands-on workflow baseline; must win side-by-side test',9 if name in ['Chaser','LeanPay','Odoo','MaintainX','Jobber','Conexiom','Productive','Eduka','DabaDoc'] else 8])

# Supplemental diligence explicitly supplies factors/fields absent from requested schemas.
supplement('16_FACTOR_SCORES','Opportunity ID|Growth Potential /10|B2B Potential /10|Targeting Ease /10|Sales Ease /10|Entry Ease /10|Gross Margin Potential /10|Demand /10|Pain /10|WTP /10|Recurring /10|CAC /10|Retention /10|High Ticket /10|Competition /10|MVP Ease /10|ROI /10|Morocco /10|Global /10|Moat /10|AI /10|Weighted Score /100|Evidence|Score Interpretation')
for o in OPPS:
    f=o['f']; p=PROFILES.get(o['id'])
    add('16_FACTOR_SCORES',[o['id'],min(9,round((f['Demand']+f['AI'])/2)),9,f['CAC'],f['CAC'],f['MVPEase'],round((p['gmlo']+p['gmhi'])/20) if p else 7,*[f[k] for k in FACTORS],o['score'],o['sources'],'All factors analyst hypotheses; Growth is heuristic, NOT measured CAGR. CAC derived from targeting/channel/sales judgment; extra factors unweighted. Higher=better, not difficulty.'])
supplement('17_SOURCES','Source ID|Publisher|URL|Publication / Period|Evidence Level|Claim and Limit|Access Date')
for s in SOURCES: add('17_SOURCES',[*s,DATE])
supplement('18_PRODUCT_DILIGENCE','Opportunity ID|10-Minute Value Test|Advanced Features|Technical Complexity /10|AI Complexity /10|Integration Complexity /10|Regulatory Complexity /10|Operational Complexity /10|Support Complexity /10|MVP Difficulty /10|Time to MVP|Retention Mechanism|Expansion|Proof Required|Confidence')
for i,p in PROFILES.items():
    add('18_PRODUCT_DILIGENCE',[i,p['demo'],p['advanced'],*p['complex'],p['weeks']+' weeks','Recurring new exceptions; history and ownership; not contractual lock-in',p['advanced'],'3 unrelated paid pilots;30–60day workflow usage and incremental ROI; buyer authority verified','Medium category pain; low wedge/WTP; very low CAC/retention. Complexity here:10=hard.'])
supplement('19_ECONOMICS_DILIGENCE','Opportunity ID|ARPU USD /month|CAC USD|Gross Margin %|Monthly Logo Churn %|Annual Logo Retention %|LTV GP USD Uncapped|LTV GP USD 24-Month Cap|LTV:CAC Capped|Payback Months|Qualified Demo-to-Paid %|Activated Trial-to-Paid %|Evidence')
for i in DEEP:
    p=PROFILES[i]; arlo=.7*p['usd']; arhi=1.5*p['usd']; gmlo=p['gmlo']/100; gmhi=p['gmhi']/100; lo=p['churnlo']/100; hi=p['churnhi']/100
    llo=arlo*gmlo/hi; lhi=arhi*gmhi/lo; caplo=arlo*gmlo*min(24,1/hi); caphi=arhi*gmhi*min(24,1/lo)
    add('19_ECONOMICS_DILIGENCE',[i,f'{arlo:.0f}–{arhi:.0f}',f'{p["caclo"]:.0f}–{p["cachi"]:.0f}',f'{p["gmlo"]:.0f}–{p["gmhi"]:.0f}',f'{p["churnlo"]}–{p["churnhi"]}',f'{100*(1-hi)**12:.1f}–{100*(1-lo)**12:.1f}',f'{llo:.0f}–{lhi:.0f}',f'{caplo:.0f}–{caphi:.0f}',f'{caplo/p["cachi"]:.1f}–{caphi/p["caclo"]:.1f}',f'{p["caclo"]/(arhi*gmhi):.1f}–{p["cachi"]/(arlo*gmlo):.1f}','10–25 (assumption, qualified demos)','15–35 (assumption, activated assisted trials)','ALL ASSUMPTIONS; international ICP, not Morocco-specific; conservative/favorable joint endpoints, not measured probability intervals. Uncapped LTV unreliable before cohorts mature.'])

supplement('20_EXPANSION_PLAYBOOK','Opportunity ID|Stage / Market|Language|Pricing|Payments|Compliance|Product|Sales|Marketing|Support|Integrations|Go / No-Go Gate')
markets=[
('1 Morocco','French admin; Arabic customer text; Darija human support','MAD; test2500 Pro for001','Bank transfer; licensed PSP only after underwriting','CNDP09-08; foreign transfers include remote support/AI; accountant reviews tax fields','Moroccan invoice IDs/payment references; no automatic legal demands','Founder + local accountant','French workflow keywords; named city/industry pages','Local business hours; defined onboarding cap','One Sage CSV; no accounting-engine replacement','3 unrelated pilots; legal data path before live processing'),
('2 Senegal OR Côte d’Ivoire','French; locally reviewed customer messages','XOF locally tested; do not simply convert MAD','Bank transfer; local licensed mobile-money/PSP partner','Country-specific privacy/tax; OHADA is not a universal software exemption','Local payment-reference and document formats','One local accounting/reseller partner','Local trade directories; one-country case study','Local partner SLA; avoid travel-dependent support','SYSCOHADA-compatible exports where relevant; chosen local ERP','2 prepaid customers; qualified partner; support economics viable'),
('3 France','French; French accounting terminology','EUR;001 Pro299 hypothesis','SEPA/card via eligible provider','GDPR/DPA/transfer assessment; approved e-invoice platform interface; avoid claiming platform approval','Dispute/e-invoice lifecycle and multi-entity mappings','Accountants + vertical CFO outbound','Evidence-led French search pages','French time-zone SLA; EU-hosting option not sufficient alone','Sage/Pennylane/Odoo via licensed APIs; verify access','2 paid pilots; one supported connector; real competitive win'),
('4 Belgium','French first; Dutch for Flanders','EUR; price-test separately','SEPA/card through eligible PSP','GDPR; local B2B e-invoice/Peppol obligations verify by scope','Belgian entity IDs/document formats','French-speaking accounting partners first','Regional FR/NL landing pages','Bilingual support if entering Flanders','Local accounting/Peppol-capable provider; do not build access point','Language and e-invoicing legal review complete'),
('5 Canada','French in Quebec; English elsewhere','CAD localized; not mislabeled USD','Bank debit/card via eligible entity/provider','PIPEDA/provincial law; Quebec Law25; CASL governs outreach','GST/HST/QST references when required; banking formats','Quebec accountants first','Quebec-specific workflow search','North American coverage before broad launch','QuickBooks/Xero; verify French settings and APIs','CAD willingness-to-pay and timezone support validated'),
('6 GCC: UAE OR Saudi Arabia','English/Arabic RTL; no assumed Darija transfer','AED or SAR; higher price requires quantified ROI','Local eligible PSP/bank transfer','UAE or Saudi privacy/data rules separately; Saudi e-invoice scope separately; counsel','Country-specific VAT/entity fields; authority matrices','Local accounting/ERP reseller','Country-specific outbound; LinkedIn only at adequate ACV','Arabic/English local-hours SLA','UAE ERP export or Saudi compliant e-invoice integration, not both at once','Choose one country; legal/PSP/data path and2 pilots'),
('7 Wider global / USA','English; other languages only with demand','USD;001 Pro349 hypothesis','Eligible card/ACH provider; tax treatment review','State privacy/security; CAN-SPAM; licensed collection boundaries; no consumer debt','US conventions; accessible product; SOC2 only if justified','Accounting channel + targeted vertical outbound','High-intent Google test after references','Time-zone coverage costed','QuickBooks/Xero first; no50-connector promise','Paid cohorts renew; acceptable CAC; scalable onboarding')]
for i in TOP15:
    p=PROFILES[i]
    for r in markets:
        row=list(r)
        if i!='SAAS-001': row[2]='Local price-test; Pro baseline '+str(p['mad'])+'MAD / '+str(p['eur'])+'EUR / '+str(p['usd'])+'USD; no assumed CAD/AED equivalents'
        row[5]+='; product-specific adaptation: '+p['features'][1]
        row[9]+='; '+p['integrations']
        add('20_EXPANSION_PLAYBOOK',[i,*row])

supplement('21_TOP20_MATRIX','Opportunity ID|Rank|SaaS Opportunity|Market|ICP|Problem|Demand|WTP|Ticket|Competition|Advertising|MVP Difficulty|Retention|Global Potential|Morocco Potential|ROI|Final Score')
for o in OPPS[:20]:
    f=o['f']; add('21_TOP20_MATRIX',[o['id'],o['rank'],o['name'],o['market'],o['icp'],o['pain'],f['Demand'],f['WTP'],price(o),f['Competition'],max(channels(o)[:3]),f['MVPEase'],f['Retention'],f['Global'],f['Morocco'],f['ROI'],o['score']])

supplement('22_METHOD_ASSUMPTIONS','Topic|Definition / Assumption|Status')
notes=[
('Research date',DATE+'; point-in-time desk research, not commercial/technical/legal due diligence','FACT'),
('Evidence','FACT/VERIFIED=primary fetched statement; STRONG SIGNAL=proxy, indexed official excerpt or selected review summary; HYPOTHESIS=untested commercial proposition; ASSUMPTION=model input; UNKNOWN=not established','METHOD'),
('Scores','Score=sum(score_i*weight_i)/10; weights in research_data.py sum100. All scores are subjective, not success probabilities. 1-point differences not decision-grade. Winner001 competition downgraded6→5 after verifying LeanPay/Upflow disputes/exception workflows; final80.6 vs00580.5 is effectively a tie.','HYPOTHESIS'),
('Spreadsheet calculations','Exports contain calculated numeric snapshots, not live spreadsheet formulas. Reproduce/recalculate with python -B build_workbook.py after changing research inputs. Formulas and assumptions are documented here.','METHOD'),
('Score direction','Requested MASTER MVP Difficulty and GLOBAL Entry Difficulty use10=EASY to match user scoring rule. Supplemental complexity and FINAL Sales Difficulty use10=HARD.','METHOD'),
('Risk gate','BUILD NOW means execute paid validation/concierge. Payroll019 and clinical039 rejected despite recurrence; CBAM033 WATCH pending specialist access. Top15 operational shortlist=001–015; numerical master ranks remain unchanged.','METHOD'),
('19 factors','Advertising separate from targeting/sales; all19 requested screening dimensions provided across master and16_FACTOR_SCORES. Additional weighted Moat and CAC included. Growth is heuristic, not empirical.','METHOD'),
('Financial FX','USD1=MAD10; EUR1=USD1.10; planning constants, not market quotations. No tax-inclusive prices.','ASSUMPTION'),
('ARPU','Pricing sheet baseline is Professional plan, not forecast blended mix. Economics diligence ARPU ranges=70–150% of proposed international Pro. Morocco must be validated separately.','ASSUMPTION'),
('LTV','Uncapped GP LTV=monthlyARPU*grossMargin/monthlyLogoChurn. Capped proxy=monthlyGP*min(24,1/churn), deliberately rough not discounted cohort DCF. Annual retention=(1-churn)^12. No NRR equivalence.','METHOD'),
('CAC','Acquisition cost includes ads, outbound tools, commissions and allocated selling labor. Never sourced empirical CAC. Model team cost excludes sales labor already in CAC.','ASSUMPTION'),
('Margins','Steady-state gross margins include hosting, usage, payment fees and delivery support. Concierge gross margin may be30–60%; one-time onboarding and implementation excluded from recurring projections.','ASSUMPTION'),
('Financial cadence','New subscriptions billed full month; churn from prior-month opening customers; fractional customer counts are cohort expected values, not observed people. MRR=end active*ARPU. ARR=12*MRR run-rate, NOT recognized annual revenue.','ASSUMPTION'),
('Financial costs','Product Costs=150 fixed/month + variable COGS. Team includes modest product-founder/engineering/support pay, excluding acquisition labor allocated to CAC. Marketing Costs=fixed content/creative only; media included in Acquisition Cost. Net Profit=pretax operating result, not statutory net income or free cash flow.','METHOD'),
('Demand/CPC','No Keyword Planner account, sampled Trends series, ad library census or Similarweb subscription. No CPC, search volume, competitor ad spend or SEO strength fabricated. Channel scores are hypotheses.','UNKNOWN'),
('Validation counts','All zero because no landing page, outreach or paid test actually executed. No stakeholder interviews conducted. Allowed decision RETEST denotes not tested; never treat zero as poor conversion.','FACT'),
('Privacy','No real customer invoices or personal data used in this research. Live pilots require lawful processing, minimization, permissions and applicable transfer formalities. A European data center alone does not solve Morocco/GDPR requirements.','METHOD'),
('Market sizing','No verified count of qualifying buyers. Require list of300 named reachable ICP accounts and pricing acceptance before bottom-up SAM. Macro company creations/tourist arrivals/WhatsApp adoption are not SaaS TAM.','UNKNOWN'),
('Competitors','50 identified competitors across10 ideas; official checks concentrated on lead comparators. All other prices, AI, sentiment/ad/SEO data remain explicit UNKNOWN. Countries are unverified origin hypotheses, not confirmed current HQ.','LIMITATION'),
('Ad economics','Allowable CAC=monthly ARPU*grossMargin*target payback months. Allowable CPC=allowable media CAC*visitor-to-qualified-lead*qualified-lead-to-paid. Derived bid ceiling is NOT observed CPC.','METHOD'),
('001 ROI','Working-capital release is not new revenue. Example MAD10m credit sales/year *5DSO days/365=MAD136986 cash release; financing saving at10% annual=MAD1142/month; labor20h*80=MAD1600/month; total2742 before fee2500, only1.10x gross benefit.','ASSUMPTION'),
('001 qualification','Prefer >=MAD20m credit sales/year or verified >=MAD7500 monthly benefit at2500 fee. Reminders cannot fix insolvent debtors or dominant buyers refusing to pay. Score moat6/10 intentionally modest.','HYPOTHESIS'),
('Founder assumption','Two full-stack engineers plus founder with French fluency and Moroccan accountant/industrial-distributor access;4–12week scopes exclude regulated system replacement. Without access, rerun ranking.','ASSUMPTION'),
('Scenario not forecast','Conservative/base/aggressive paths are independent assumptions. No TAM-derived customer forecast or guaranteed acquisition capacity. Fund validation in stages.','ASSUMPTION')]
for row in notes: add('22_METHOD_ASSUMPTIONS',list(row))

stages=[('1–10',10,200,.05,500,'Founder + accountant introductions','30-day paid pilot atMAD2500; onboarding capped','1–3 months','Contact owners/CFOs of credit-selling industrial wholesalers; show exception audit'),('11–25',15,300,.05,700,'Partner referrals + targeted outbound','Same price; reference-backed30day pilot','2–4 additional months','Recruit3 active accountants; publish1 consented case study'),('26–50',25,500,.05,850,'Outbound + partners; small exact-match Search test','Paid pilot; no open-ended free service','3–5 additional months','Measure channel CAC; introduce implementation fee for complex imports'),('51–100',50,1000,.05,1000,'Repeatable partner channel + Search/retargeting','Standard onboarding; annual option after value proof','4–8 additional months','Hire only after documented conversion; track weekly cohorts')]
for s,n,leads,conv,cac,ch,offer,tm,action in stages:
    add('12_FIRST_100_CUSTOMERS',['SAAS-001',s,s,'ASSUMPTION; targets, not promised net customer counts',ch,BYID['SAAS-001']['icp'],offer,leads,f'{100*conv:.0f}% account-to-paid assumption',n,f'USD{cac}',f'USD{n*250} incremental gross MRR before churn',tm,action])

SCENARIOS={
'Conservative':dict(new=[0,0,1,1,1,2,2,2,2,3,3,3]+[3]*12,arpu=200,churn=.025,cac=1000,gm=.72,team=[2500]*6+[3500]*6+[5000]*12,marketing=100,other=300),
'Base':dict(new=[0,2,3,4,5,6,7,8,8,9,10,10]+[12]*12,arpu=250,churn=.015,cac=800,gm=.82,team=[3000]*3+[4500]*3+[6500]*6+[10000]*12,marketing=300,other=500),
'Aggressive':dict(new=[1,3,5,7,10,12,15,18,20,22,25,28]+[30]*12,arpu=300,churn=.01,cac=650,gm=.85,team=[4500]*3+[7000]*3+[12000]*6+[22000]*12,marketing=600,other=900)}
summary=[]
for scenario,p in SCENARIOS.items():
    cust=0; cumulative=0; trough=0; recs=[]
    for m,new in enumerate(p['new'],1):
        churned=cust*p['churn']; cust=cust-churned+new; mrr=cust*p['arpu']; acq=new*p['cac']; product=150+mrr*(1-p['gm']); team=p['team'][m-1]; total=acq+product+team+p['marketing']+p['other']; gp=mrr-product; net=mrr-total
        row=['SAAS-001',scenario,m,cust,new,churned,p['arpu'],mrr,12*mrr,p['cac'],acq,product,team,p['marketing'],p['other'],total,gp,net]
        add('13_FINANCIAL_MODEL',[round(v,2) if isinstance(v,float) else v for v in row]); recs.append(row); cumulative+=net; trough=min(trough,cumulative)
    for start,end,label in [(1,3,'Months1–3'),(4,6,'Months4–6'),(7,12,'Months7–12'),(13,24,'Year2')]:
        rs=recs[start-1:end]; last=rs[-1]
        summary.append([scenario,label,round(last[3],1),sum(r[4] for r in rs),round(sum(r[5] for r in rs),1),p['arpu'],round(last[7]),round(last[8]),p['cac'],round(sum(r[10] for r in rs)),round(sum(r[13] for r in rs)),round(sum(r[11]+r[12]+r[14] for r in rs)),round(sum(r[16] for r in rs)),round(sum(r[17] for r in rs)),round(-trough)])
supplement('23_SCENARIO_SUMMARY','Scenario|Period|End Customers|New Customers Period|Churned Customers Period|ARPU USD|Exit MRR USD|Exit ARR Run-rate USD|CAC USD|Acquisition Spend Period USD|Fixed Marketing Period USD|COGS Team Other Period USD|Gross Profit Period USD|Operating Profit Period USD|Peak24-Month Cumulative Deficit USD')
for r in summary:add('23_SCENARIO_SUMMARY',r)

wb=Workbook(); wb.remove(wb.active)
for idx,(name,rows) in enumerate(sheets.items(),1):
    ws=wb.create_sheet(name); ws.freeze_panes='C2' if rows[0][0]=='Opportunity ID' else 'A2'
    for row in rows: ws.append(row)
    for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='12344D'); cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=44
    for col in range(1,len(rows[0])+1): ws.column_dimensions[get_column_letter(col)].width=min(54,max(17,len(str(rows[0][col-1]))+3))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(cell.value,float): cell.number_format='0.00'
    if len(rows)>1:
        tab=Table(displayName='Research'+str(idx),ref=f'A1:{get_column_letter(len(rows[0]))}{len(rows)}'); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(tab)
    with (CSV/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f: csv.writer(f).writerows(rows)
path=OUT/'SaaS_Opportunity_Research_2026-09-15.xlsx'; wb.save(path)
# Cross-reference, numerical and schema checks use reopened artifact, not only input state.
check=load_workbook(path,data_only=True)
for name,headers in SCHEMAS.items(): assert [x.value for x in check[name][1]]==headers
assert len(check['01_MASTER_RANKING']['A'])==41
assert len(sheets['08_COMPETITORS'])==51
assert len(sheets['13_FINANCIAL_MODEL'])==73
for name,rows in sheets.items():
    if rows[0][0]=='Opportunity ID': assert all(r[0] in BYID for r in rows[1:]),name
for r in sheets['13_FINANCIAL_MODEL'][1:]:
    assert abs(r[7]-r[3]*r[6])<2
    assert abs(r[17]-(r[7]-r[15]))<.03
    assert abs(r[15]-sum(r[10:15]))<.04
with (OUT/'validation_results.json').open('w') as f: json.dump({'date':DATE,'schema_checks':15,'opportunities':40,'detailed_products':15,'competitor_rows':50,'pricing_rows':315,'scenario_months':72,'sheets':{k:len(v)-1 for k,v in sheets.items()},'all_checks_passed':True},f,indent=2)
with zipfile.ZipFile(OUT/'Google_Sheets_CSV_2026-09-15.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p in sorted(CSV.glob('*.csv')): z.write(p,'csv/'+p.name)
print('Workbook:',path)
print('Scenario summary:',json.dumps(summary,indent=2))
print('All schema and financial checks passed.')
